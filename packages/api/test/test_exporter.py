import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from egapro import csv_representation, db, dgt, dgt_representation, exporter


@pytest.mark.parametrize(
    "module", [exporter, dgt, dgt_representation, csv_representation]
)
def test_clean_cell_neutralizes_formulas(module):
    assert module.clean_cell('=HYPERLINK("http://evil.com")') == '\'=HYPERLINK("http://evil.com")'
    # Leading whitespaces are stripped before checking.
    assert module.clean_cell("  \t=SUM(A1)") == "'=SUM(A1)"
    assert module.clean_cell("Total  Recall") == "Total Recall"
    assert module.clean_cell(-3) == -3
    assert module.clean_cell(12.5) == 12.5
    assert module.clean_cell(None) is None


@pytest.mark.parametrize(
    "module", [exporter, dgt, dgt_representation, csv_representation]
)
def test_clean_cell_keeps_xlsx_text_readable(module):
    # openpyxl writes those as text cells, never evaluated: no visible quote.
    assert module.clean_cell("+33 6 12 34 56 78") == "+33 6 12 34 56 78"
    assert module.clean_cell("-- Non renseigné --") == "-- Non renseigné --"
    assert module.clean_cell("@SUM(A1)") == "@SUM(A1)"


@pytest.mark.parametrize(
    "module", [exporter, dgt, dgt_representation, csv_representation]
)
def test_clean_cell_writes_no_formula_in_xlsx(module):
    values = ["=1+1", "  =cmd|' /C calc'!A0", "+1+1", "-1+1", "@SUM(1)", "\t=1", "\r=1"]
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append([module.clean_cell(v) for v in values])
    out = io.BytesIO()
    wb.save(out)
    sheet = zipfile.ZipFile(out).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet
    assert [c.value for c in next(load_workbook(out).active.iter_rows())] == [
        "'=1+1",
        "'=cmd|' /C calc'!A0",
        "+1+1",
        "-1+1",
        "@SUM(1)",
        "'=1",
        "'=1",
    ]


@pytest.fixture
async def init_db():
    await db.init()
    yield
    await db.terminate()


@pytest.mark.asyncio
async def test_public_data_csv_neutralizes_formulas(init_db, declaration):
    await declaration(siren="514027945", year=2020, company="=1+1", grade=26)
    out = io.StringIO()
    await exporter.public_data(out)
    lines = out.getvalue().splitlines()
    assert len(lines) == 2
    cells = lines[1].split(";")
    assert cells[3] == "514027945"
    assert cells[4] == "'=1+1"
    assert cells[-1] == "26"
